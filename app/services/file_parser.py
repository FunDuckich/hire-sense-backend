import io
import asyncio
import traceback
from docx import Document
import fitz
from fastapi import UploadFile, HTTPException
from striprtf.striprtf import rtf_to_text
from app.services.ocr import ocr_service


def _parse_docx_sync(file_bytes: bytes) -> str:
    try:
        doc = Document(io.BytesIO(file_bytes))
        full_text = []

        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())

        def extract_from_table(table):
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if para.text.strip():
                            full_text.append(para.text.strip())
                    for nested_table in cell.tables:
                        extract_from_table(nested_table)

        for table in doc.tables:
            extract_from_table(table)

        return "\n\n".join(full_text)

    except Exception as e:
        raise RuntimeError(f"Error parsing DOCX content: {e}") from e


async def parse_file(file: UploadFile) -> str:
    content_type = file.content_type
    file_content = await file.read()

    try:
        text = ""
        if content_type == "application/pdf":
            pdf_document = fitz.open(stream=file_content, filetype="pdf")
            for page in pdf_document:
                text += page.get_text()
            pdf_document.close()

            if not text.strip():
                print("Текстовый слой пуст, запускаем OCR-фолбэк...")
                image_bytes_list = []
                pdf_for_ocr = fitz.open(stream=io.BytesIO(file_content), filetype="pdf")
                for page in pdf_for_ocr:
                    pix = page.get_pixmap(dpi=300)
                    image_bytes_list.append(pix.tobytes("png"))
                pdf_for_ocr.close()

                loop = asyncio.get_running_loop()
                text = await loop.run_in_executor(None, ocr_service.recognize, image_bytes_list)

        elif content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            loop = asyncio.get_running_loop()
            text = await loop.run_in_executor(None, _parse_docx_sync, file_content)

        elif content_type in ["application/rtf", "text/rtf"]:
            try:
                decoded_content = file_content.decode('cp1251')
            except UnicodeDecodeError:
                decoded_content = file_content.decode('utf-8', errors='ignore')
            text = rtf_to_text(decoded_content)

        else:
            raise HTTPException(status_code=415, detail="Unsupported file type")

        if not text or not text.strip():
            raise HTTPException(status_code=400,
                                detail="Не удалось извлечь текст из файла. Файл может быть пустым или поврежденным.")

        return text.strip()

    except HTTPException as e:
        raise e
    except Exception as e:
        print(f"UNEXPECTED PARSING ERROR: {e}")
        traceback.print_exc()
        raise HTTPException(status_code=500, detail="Внутренняя ошибка при обработке файла.")
